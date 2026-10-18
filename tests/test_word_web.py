import re
from pathlib import Path

from docx import Document

from dawaa import outils
from dawaa.llm.agent import MAX_CARACTERES_HISTORIQUE, preparer_historique
from dawaa.outils import web, word


def test_creer_puis_lire_word():
    resultat = word.creer_document_word(
        "Lettre de réclamation",
        "Madame, Monsieur,\n\n## Objet\n- premier point\n- **point important**\n\nCordialement.",
        nom_fichier="Réclamation",
    )
    chemin = Path(re.search(r"ne pas montrer\) : (.+?\.docx)", resultat).group(1))
    assert chemin.name == "Réclamation.docx"
    doc = Document(chemin)
    assert doc.paragraphs[0].text == "Lettre de réclamation"
    gras = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "point important" in gras
    assert "Cordialement." in word.lire_document(str(chemin))


def test_interface_retrouve_le_fichier_cree():
    from dawaa.app import FICHIER_RESULTAT

    resultat = word.creer_document_word("Note", "Bonjour")
    assert "Emplacement à annoncer à l'utilisateur : " in resultat
    (chemin,) = FICHIER_RESULTAT.findall(resultat)
    assert Path(chemin).is_file()


def test_lire_document_tronque(tmp_path):
    f = tmp_path / "long.txt"
    f.write_text("a" * 10_000, encoding="utf-8")
    assert "document tronqué" in word.lire_document(str(f))


def test_recherche_web_simulee(monkeypatch):
    class FauxDDGS:
        def text(self, *a, **k):
            return [{"title": "Météo Lyon", "body": "Soleil, 18 °C", "href": "https://meteo.example"}]
    monkeypatch.setattr(web, "DDGS", FauxDDGS)
    texte = outils.executer("recherche_web", '{"requete": "météo Lyon"}')
    assert "Soleil" in texte and "https://meteo.example" in texte


def test_recherche_web_sans_internet(monkeypatch):
    class FauxDDGS:
        def text(self, *a, **k):
            raise RuntimeError("Connection refused")
    monkeypatch.setattr(web, "DDGS", FauxDDGS)
    assert "Pas de connexion internet" in outils.executer("recherche_web", {"requete": "x"})


def test_outil_inconnu_et_json_invalide():
    assert "n'existe pas" in outils.executer("formater_disque", {})
    assert "illisibles" in outils.executer("recherche_web", "{pas du json")


def test_historique_raccourci_commence_par_user():
    messages = []
    for i in range(40):
        messages += [{"role": "user", "content": f"question {i} " + "x" * 300},
                     {"role": "assistant", "content": "réponse " + "y" * 300}]
    reduit = preparer_historique(messages)
    assert reduit[0]["role"] == "user"
    assert sum(len(m["content"]) for m in reduit) <= MAX_CARACTERES_HISTORIQUE + 100
    assert reduit[-1] == messages[-1]


def test_lire_odt_libreoffice():
    """exemple.odt a été produit par LibreOffice : titre, gras, liste, tableau, espaces multiples."""
    texte = word.lire_document(str(Path(__file__).parent / "exemple.odt"))
    assert "Contrat de location" in texte
    assert "Le loyer mensuel est de 650 euros, charges comprises." in texte
    assert "- Premier point de la liste" in texte
    assert "Poste | Montant" in texte and "Loyer | 650" in texte
    assert "Fait à Lyon,    le 18 octobre." in texte


def test_odt_abime(tmp_path):
    f = tmp_path / "abime.odt"
    f.write_bytes(b"ceci n'est pas une archive")
    assert outils.executer("lire_document", {"chemin": str(f)}).startswith("Erreur : Je n'arrive pas à lire ce document")


def test_odt_accepte_par_l_interface():
    from dawaa.app import EXTENSIONS_JOINTES
    assert ".odt" in EXTENSIONS_JOINTES
