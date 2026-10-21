import pytest
from docx import Document
from openpyxl import Workbook

from raya import outils
from raya.outils import fichiers, navigation, ordinateur, word
from raya.outils.fichiers import ErreurOutil


@pytest.fixture
def bureau():
    return fichiers.dossiers_personnels()["bureau"]


@pytest.fixture
def ouverts(monkeypatch):
    liste = []
    monkeypatch.setattr(ordinateur, "ouvrir", liste.append)
    return liste


def _word(chemin, texte):
    doc = Document()
    doc.add_paragraph(texte)
    doc.save(chemin)


def test_lire_un_fichier_du_bureau_par_une_partie_du_nom(bureau):
    _word(bureau / "Factures EDF 2026.docx", "Montant : 84 euros")
    assert "84 euros" in word.lire_document("factures edf", dossier="mon bureau")
    assert "84 euros" in word.lire_document("Factures EDF")  # sans préciser le dossier


def test_lire_un_sous_dossier_et_un_texte_dicte(bureau):
    (bureau / "Papiers").mkdir()
    (bureau / "Papiers" / "notes.txt").write_text("Rendez-vous mardi", encoding="utf-8")
    assert "mardi" in word.lire_document("notes point txt", dossier="Bureau")


def test_lire_un_tableau_excel_du_bureau(bureau):
    wb = Workbook()
    wb.active.append(["Mois", "Dépense"])
    wb.active.append(["Janvier", 120])
    wb.save(bureau / "budget.xlsx")
    assert "Janvier" in outils.executer("lire_tableau_excel", {"chemin": "budget", "dossier": "Bureau"})


def test_mauvais_outil_redirige_vers_le_bon(bureau):
    Workbook().save(bureau / "budget.xlsx")
    resultat = outils.executer("lire_document", {"chemin": "budget"})
    assert "tableau Excel" in resultat and "outil tableau, action=lire" in resultat


def test_nom_ambigu_liste_les_possibilites(bureau):
    documents = fichiers.dossiers_personnels()["documents"]
    _word(bureau / "lettre.docx", "a")
    _word(documents / "lettre.docx", "b")
    with pytest.raises(ErreurOutil, match="Plusieurs fichiers") as e:
        word.lire_document("lettre")
    assert str(e.value).count("lettre.docx (chemin pour les outils") == 2
    # Le nom exact l'emporte sur un nom qui le contient seulement
    _word(bureau / "lettre au maire.docx", "Monsieur le maire")
    assert "Monsieur le maire" in word.lire_document("lettre au maire")


def test_introuvable_liste_le_contenu_du_dossier(bureau):
    (bureau / "photo vacances.jpg").write_bytes(b"x")
    with pytest.raises(ErreurOutil, match="photo vacances.jpg"):
        word.lire_document("contrat", dossier="Bureau")


def test_dossier_inconnu():
    assert "Dossiers possibles" in outils.executer("lire_document", {"chemin": "x", "dossier": "grenier"})


def test_fichiers_caches_ignores(bureau):
    (bureau / "~$lettre.docx").write_bytes(b"verrou Word")
    (bureau / "desktop.ini").write_text("x")
    with pytest.raises(ErreurOutil, match="Je ne trouve pas"):
        fichiers.chercher("lettre")


def test_ouvrir_un_fichier(bureau, ouverts):
    (bureau / "Recette crêpes.pdf").write_bytes(b"%PDF")
    resultat = outils.executer("ouvrir_fichier", {"nom": "recette crepes", "dossier": "sur le bureau"})
    assert ouverts == [bureau / "Recette crêpes.pdf"]
    assert "Fichier ouvert" in resultat


def test_ouvrir_un_dossier(bureau, ouverts):
    (bureau / "Photos de Noël").mkdir()
    outils.executer("ouvrir_fichier", {"nom": "photos de noel"})
    outils.executer("ouvrir_fichier", {"nom": "mes téléchargements"})
    outils.executer("ouvrir_fichier", {"nom": "mes téléchargements", "dossier": "Téléchargements"})
    outils.executer("ouvrir_fichier", {"dossier": "Bureau"})
    assert ouverts == [bureau / "Photos de Noël", fichiers.dossiers_personnels()["telechargements"],
                       fichiers.dossiers_personnels()["telechargements"], bureau]


def test_jamais_de_programme(bureau, ouverts):
    (bureau / "Chrome.lnk").write_bytes(b"x")
    (bureau / "installer.exe").write_bytes(b"x")
    assert "je n'ouvre pas les programmes" in outils.executer("ouvrir_fichier", {"nom": "Chrome"})
    assert "je n'ouvre pas les programmes" in outils.executer("ouvrir_fichier", {"nom": str(bureau / "installer.exe")})
    assert ouverts == []


def test_vrai_bureau_windows():
    """Le Bureau donné par Windows existe (il peut être dans OneDrive)."""
    _, anglais, identifiant = fichiers.DOSSIERS_WINDOWS["bureau"]
    assert fichiers._dossier_windows(identifiant, anglais).is_dir()


def test_ouvrir_quand_on_demande_de_lire(bureau, ouverts, monkeypatch):
    """Le modèle appelle parfois ouvrir_fichier pour « lis-moi » : on lit au lieu d'ouvrir."""
    (bureau / "courses.txt").write_text("lait, pain", encoding="utf-8")
    monkeypatch.setattr(navigation.contexte, "derniere_demande", "Lis-moi ma liste de courses sur le bureau")
    assert "lait, pain" in outils.executer("ouvrir_fichier", {"nom": "courses", "dossier": "Bureau"})
    assert ouverts == []
    monkeypatch.setattr(navigation.contexte, "derniere_demande", "Ouvre et lis-moi mes courses")
    assert "lait, pain" in outils.executer("ouvrir_fichier", {"nom": "courses"})
    assert ouverts == [bureau / "courses.txt"]


def test_emplacement_en_francais(bureau):
    assert fichiers.emplacement_lisible(bureau / "Papiers" / "facture.pdf") == "Bureau > Papiers > facture.pdf"
    assert fichiers.emplacement_lisible(fichiers.dossiers_personnels()["images"]) == "Images"


def test_dossier_devine_a_tort(bureau):
    """Le modèle invente « Documents » alors que le fichier est sur le Bureau : on le retrouve quand même."""
    _word(bureau / "Courrier mairie.docx", "réunion jeudi")
    assert "réunion jeudi" in word.lire_document("courrier mairie", dossier="Documents")


def test_dossier_non_cite_par_l_utilisateur_ignore(bureau, monkeypatch):
    """« Lis-moi ma liste de courses » : le modèle ajoute dossier=Documents, on cherche partout et on demande lequel."""
    documents = fichiers.dossiers_personnels()["documents"]
    (bureau / "Liste de courses.txt").write_text("lait", encoding="utf-8")
    (documents / "Liste de courses.txt").write_text("riz", encoding="utf-8")
    monkeypatch.setattr(navigation.contexte, "derniere_demande", "Lis-moi ma liste de courses")
    assert "Plusieurs fichiers" in outils.executer("lire_document", {"chemin": "liste de courses", "dossier": "Documents"})
    monkeypatch.setattr(navigation.contexte, "derniere_demande", "Lis la liste de courses de mes documents")
    assert "riz" in outils.executer("lire_document", {"chemin": "liste de courses", "dossier": "Documents"})


def test_relance_si_le_modele_demande_ou_est_le_fichier():
    from raya.llm.agent import CORRECTION_RECHERCHE, correction_necessaire
    reponse = "Pouvez-vous me donner le nom du fichier ou me dire où il se trouve ?"
    assert correction_necessaire("Lis-moi ma liste de courses", reponse) == CORRECTION_RECHERCHE
    assert correction_necessaire("Bonjour", "Bonjour ! Comment allez-vous ?") is None


def test_petits_mots_ignores(bureau):
    _word(bureau / "Courrier mairie.docx", "réunion jeudi")
    assert "réunion jeudi" in word.lire_document("le courrier de la mairie")
    assert "réunion jeudi" in word.lire_document("fichier courrier mairie")
