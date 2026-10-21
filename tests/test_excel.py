import re
from pathlib import Path

from openpyxl import load_workbook

from raya import outils
from raya.outils import excel


def _chemin(resultat: str) -> Path:
    return Path(re.search(r"ne pas montrer\) : (.+?\.xlsx)", resultat).group(1))


def _creer():
    return _chemin(excel.creer_tableau_excel(
        "Dépenses", ["Poste", "Montant", "Quantité"],
        [["Loyer", "650", 1], ["Courses", "320,50 €", 4], ["Essence", 80, 2]],
        ligne_total=True,
    ))


def test_creation_avec_total():
    chemin = _creer()
    ws = load_workbook(chemin).active
    assert [c.value for c in ws[1]] == ["Poste", "Montant", "Quantité"]
    assert ws["B3"].value == 320.5  # « 320,50 € » converti en nombre
    assert ws["A5"].value == "Total"
    assert ws["B5"].value == "=SUM(B2:B4)"
    assert ws["A1"].font.bold


def test_jamais_ecraser_a_la_creation():
    a, b = _creer(), _creer()
    assert a != b and a.exists() and b.exists()
    assert b.name == "Dépenses (2).xlsx"


def test_lecture_calcule_les_statistiques():
    texte = excel.lire_tableau_excel(str(_creer()))
    assert "3 ligne(s) de données" in texte
    assert "Montant : total = 1050.5" in texte
    assert "max = 650" in texte


def test_lecture_par_simple_nom():
    _creer()
    assert "Poste | Montant" in excel.lire_tableau_excel("Dépenses")


def test_modifications_sur_une_copie():
    original = _creer()
    resultat = excel.modifier_tableau_excel(str(original), [
        {"type": "ajouter_ligne", "valeurs": ["Internet", 30, 1]},
        {"type": "ajouter_colonne", "nom": "Coût total", "operation": "multiplier", "colonnes": ["Montant", "Quantité"]},
        {"type": "trier", "colonne": "montant", "ordre": "croissant"},
    ])
    copie = _chemin(resultat)
    assert copie.name == "Dépenses (modifié).xlsx"
    assert load_workbook(original).active.max_row == 5  # original intact
    ws = load_workbook(copie).active
    assert [ws.cell(row=r, column=1).value for r in range(2, 6)] == ["Internet", "Essence", "Courses", "Loyer"]
    assert ws["A6"].value == "Total" and ws["B6"].value == "=SUM(B2:B5)"
    # La formule a suivi sa ligne lors du tri
    assert ws["D1"].value == "Coût total" and ws["D2"].value == "=B2*C2" and ws["D5"].value == "=B5*C5"


def test_modification_successive_ne_ecrase_pas():
    original = _creer()
    c1 = _chemin(excel.modifier_tableau_excel(str(original), [{"type": "mettre_en_forme"}]))
    c2 = _chemin(excel.modifier_tableau_excel(str(c1), [{"type": "mettre_en_forme"}]))
    assert c2.name == "Dépenses (modifié 2).xlsx"


def test_colonne_inconnue_renvoie_un_message_clair():
    resultat = outils.executer("tableau", {
        "action": "modifier", "chemin": str(_creer()), "modifications": [{"type": "trier", "colonne": "Prix"}]})
    assert resultat.startswith("Erreur : Il n'y a pas de colonne « Prix »")
    assert "Montant" in resultat


def test_graphique():
    copie = _chemin(excel.ajouter_graphique_excel(str(_creer()), "barres", "Poste", ["Montant"], "Mes dépenses"))
    assert len(load_workbook(copie).active._charts) == 1


def test_fichier_absent():
    assert "Je ne trouve pas" in outils.executer("tableau", {"action": "lire", "chemin": "inexistant.xlsx"})


# ---------- un seul outil pour les tableaux ----------

def test_outil_tableau_et_ses_actions():
    cree = outils.executer("tableau", {"action": "creer", "nom": "Courses", "colonnes": ["Article", "Prix"],
                                       "lignes": [["Pain", 1.2], ["Lait", 0.95]], "ligne_total": True})
    chemin = str(_chemin(cree))
    assert "Prix : total = 2.15" in outils.executer("tableau", {"action": "lire", "chemin": "Courses"})
    modifie = outils.executer("tableau", {"action": "modifier", "chemin": chemin,
                                          "modifications": [{"type": "ajouter_ligne", "valeurs": ["Oeufs", 2.5]}]})
    assert "Modifications faites" in modifie
    graphique = outils.executer("tableau", {"action": "graphique", "chemin": chemin, "graphique": "camembert",
                                            "colonne_x": "Article", "colonnes_y": ["Prix"]})
    assert len(load_workbook(_chemin(graphique)).active._charts) == 1
    assert "Action inconnue" in outils.executer("tableau", {"action": "supprimer"})


def test_action_oubliee_devinee():
    """Le petit modèle oublie parfois action : on la devine d'après les paramètres."""
    chemin = str(_chemin(outils.executer("tableau", {"nom": "Notes", "colonnes": ["Élève", "Note"],
                                                     "lignes": [["Léa", 15]]})))
    assert "Feuilles :" in outils.executer("tableau", {"chemin": chemin})
    assert "Graphique ajouté" in outils.executer("tableau", {"chemin": chemin, "colonne_x": "Élève", "colonnes_y": ["Note"]})


def test_anciens_noms_d_outils_acceptes():
    """Les raccourcis enregistrés avant la fusion appellent encore lire_tableau_excel."""
    chemin = str(_creer())
    assert "Feuilles :" in outils.executer("lire_tableau_excel", {"chemin": chemin})
    assert "Modifications faites" in outils.executer("modifier_tableau_excel", {
        "chemin": chemin, "actions": [{"type": "mettre_en_forme"}]})
    assert "Graphique ajouté" in outils.executer("ajouter_graphique_excel", {
        "chemin": chemin, "type": "barres", "colonne_x": "Poste", "colonnes_y": ["Montant"]})


def test_lecture_de_tableau_rejouable_dans_un_raccourci():
    chemin = str(_creer())
    outils.executer("tableau", {"action": "creer", "nom": "Autre", "colonnes": ["A"], "lignes": []})
    outils.executer("tableau", {"action": "lire", "chemin": chemin})
    assert [(a["outil"], a["arguments"]["action"]) for a in outils.journal] == [("tableau", "lire")]


def test_message_d_attente_selon_l_action():
    assert outils.message_attente("tableau", '{"action": "graphique", "chemin": "x"}') == "J'ajoute le graphique…"
    assert outils.message_attente("tableau", '{"action":"lire"}') == "Je lis le tableau…"
    assert outils.message_attente("lire_tableau_excel", "{}") == "Je travaille sur le tableau…"
    assert outils.message_attente("memoire", '{"action": "lancer_raccourci"}') == "Je lance le raccourci…"


def test_modifications_mal_formees_comprises():
    """Le petit modèle écrit {"ajouter_ligne": {...}} ou « tri » au lieu de la forme attendue."""
    resultat = outils.executer("tableau", {"action": "modifier", "chemin": str(_creer()), "modifications": [
        {"ajouter_ligne": {"valeurs": ["Électricité", 95]}}, {"type": "tri", "colonne": "Montant", "ordre": "decroissant"}]})
    assert "ligne ajoutée" in resultat and "trié par « Montant »" in resultat
    lignes = list(load_workbook(_chemin(resultat)).active.iter_rows(min_row=2, max_row=5, max_col=1, values_only=True))
    assert [l[0] for l in lignes] == ["Loyer", "Courses", "Électricité", "Essence"]
    erreur = outils.executer("tableau", {"action": "modifier", "chemin": str(_creer()), "modifications": [{"colorier": {}}]})
    assert "Modification inconnue" in erreur and '"type": "ajouter_ligne"' in erreur


def test_colonne_calculee_avec_un_nombre():
    """« Ajoute une colonne Par mois qui divise le montant par 12 »."""
    for modification in ({"operation": "diviser", "colonnes": ["Montant"], "nombre": 12},
                         {"operation": "diviser", "colonnes": ["Montant", "12"]}):
        resultat = outils.executer("tableau", {"action": "modifier", "chemin": str(_creer()), "modifications": [
            {"type": "ajouter_colonne", "nom": "Par mois", **modification}]})
        ws = load_workbook(_chemin(resultat)).active
        assert ws["D1"].value == "Par mois" and ws["D2"].value == "=B2/12"
    erreur = outils.executer("tableau", {"action": "modifier", "chemin": str(_creer()), "modifications": [
        {"type": "ajouter_colonne", "nom": "Par mois", "operation": "diviser", "colonnes": ["Montant"]}]})
    assert '"nombre": 12' in erreur


def test_supprimer_une_colonne():
    resultat = outils.executer("tableau", {"action": "modifier", "chemin": str(_creer()),
                                           "modifications": [{"type": "effacer_colonne", "colonne": "Quantité"}]})
    assert "colonne « Quantité » supprimée" in resultat
    assert [c.value for c in load_workbook(_chemin(resultat)).active[1]] == ["Poste", "Montant"]


def test_refus_de_modifier_un_tableau_corrige():
    from raya.llm.agent import CORRECTION_TABLEAU, correction_necessaire
    reponse = "Je ne peux pas supprimer une ligne de votre fichier Excel directement. Vous devez le faire vous-même."
    assert correction_necessaire("Efface la première ligne.", reponse) == CORRECTION_TABLEAU
    assert correction_necessaire("Supprime le fichier budget", "Je ne peux pas supprimer de fichier.") is None
