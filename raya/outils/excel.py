"""Outils Excel : créer, lire, modifier un tableau et y ajouter un graphique.

Tous les calculs sont faits ici, en Python : le modèle ne calcule jamais « de tête ».
Les anciens fichiers Excel (.xls), LibreOffice Calc (.ods) et CSV sont convertis en mémoire ;
une modification produit alors une copie au format .xlsx.
"""
import csv
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from raya.stockage.memoire import normaliser
from raya.outils.fichiers import (
    ErreurOutil, copie_modifiee, decrire_fichier, emplacement_lisible, fichier_existant, nouveau_fichier, ouvrir,
)

EXTENSIONS = (".xlsx", ".xlsm")
EXTENSIONS_LECTURE = (*EXTENSIONS, ".xls", ".ods", ".csv")
MAX_LIGNES_CONVERSION = 20_000
MAX_COLONNES_CONVERSION = 200
TAILLE_MAX_CONVERSION = 50 * 1024 * 1024
LIGNES_APERCU = 15
REMPLISSAGE_ENTETE = PatternFill("solid", fgColor="DCE6F1")


# ---------- aides ----------

def _nombre(valeur):
    """« 12,50 € » → 12.5 ; laisse le texte tel quel."""
    if isinstance(valeur, (int, float)) or valeur is None:
        return valeur
    texte = str(valeur).strip()
    nettoye = texte.replace("€", "").replace("%", "").replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        n = float(nettoye)
    except ValueError:
        return texte
    return int(n) if n.is_integer() and "." not in nettoye else n


def _mettre_en_forme(ws: Worksheet) -> None:
    if ws.max_row < 1:
        return
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
        cellule.fill = REMPLISSAGE_ENTETE
    for i, colonne in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        largeur = max((len(str(c.value)) for c in colonne if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(largeur + 2, 10), 60)
    ws.freeze_panes = "A2"
    derniere = ws.max_row - (1 if _a_ligne_total(ws) else 0)
    if derniere >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{derniere}"


def _entetes(ws: Worksheet) -> list[str]:
    return [str(c.value).strip() if c.value is not None else "" for c in ws[1]]


def _index_colonne(ws: Worksheet, nom: str) -> int:
    """Numéro (1, 2…) de la colonne nommée « nom », sans tenir compte des majuscules."""
    entetes = [e.lower() for e in _entetes(ws)]
    cle = str(nom).strip().lower()
    if cle in entetes:
        return entetes.index(cle) + 1
    for i, e in enumerate(entetes):  # correspondance partielle : « montant » ↔ « Montant (€) »
        if cle and cle in e:
            return i + 1
    if not any(entetes):
        raise ErreurOutil("Ce tableau ne commence pas tout en haut de la feuille : je ne sais modifier que les tableaux "
                          "dont les titres de colonnes sont sur la ligne 1. Je peux seulement le lire.")
    raise ErreurOutil(f"Il n'y a pas de colonne « {nom} ». Colonnes disponibles : {', '.join(_entetes(ws))}.")


def _a_ligne_total(ws: Worksheet) -> bool:
    premiere = ws.cell(row=ws.max_row, column=1).value
    return ws.max_row > 1 and isinstance(premiere, str) and premiere.strip().lower() == "total"


def _derniere_ligne_donnees(ws: Worksheet) -> int:
    return ws.max_row - 1 if _a_ligne_total(ws) else ws.max_row


def _colonnes_numeriques(ws: Worksheet) -> list[int]:
    resultat = []
    for col in range(1, ws.max_column + 1):
        valeurs = [ws.cell(row=r, column=col).value for r in range(2, _derniere_ligne_donnees(ws) + 1)]
        remplies = [v for v in valeurs if v is not None]
        if remplies and all(isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("=")) for v in remplies) \
                and any(isinstance(v, (int, float)) for v in remplies):
            resultat.append(col)
    return resultat


def _ecrire_ligne_total(ws: Worksheet, colonnes: list[int] | None = None) -> None:
    if _a_ligne_total(ws):
        ws.delete_rows(ws.max_row)
    derniere = ws.max_row
    ligne = derniere + 1
    ws.cell(row=ligne, column=1, value="Total").font = Font(bold=True)
    for col in colonnes or _colonnes_numeriques(ws):
        if col == 1:
            continue
        lettre = get_column_letter(col)
        c = ws.cell(row=ligne, column=col, value=f"=SUM({lettre}2:{lettre}{derniere})")
        c.font = Font(bold=True)


# ---------- autres formats : .xls, .ods, .csv ----------

def _sans_vides(lignes: list[list]) -> list[list]:
    """Retire les cellules vides en fin de ligne et les lignes vides en fin de feuille."""
    lignes = [l[:max((i + 1 for i, v in enumerate(l) if v is not None), default=0)] for l in lignes]
    while lignes and not lignes[-1]:
        lignes.pop()
    return lignes


def _feuilles_xls(source: Path) -> list[tuple[str, list[list]]]:
    classeur = xlrd.open_workbook(source, on_demand=True)
    feuilles = []
    for feuille in classeur.sheets():
        lignes = []
        for r in range(min(feuille.nrows, MAX_LIGNES_CONVERSION)):
            ligne = []
            for cellule in feuille.row(r)[:MAX_COLONNES_CONVERSION]:
                v = cellule.value
                if cellule.ctype == xlrd.XL_CELL_NUMBER:
                    v = int(v) if float(v).is_integer() else v
                elif cellule.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate_as_datetime(v, classeur.datemode)
                    except (ValueError, OverflowError):
                        pass
                elif cellule.ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif cellule.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    v = None
                ligne.append(v)
            lignes.append(ligne)
        feuilles.append((feuille.name, _sans_vides(lignes)))
    return feuilles


ODS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


def _attribut(element, espace: str, nom: str, defaut=None):
    return element.get(f"{{{ODS[espace]}}}{nom}", defaut)


def _valeur_ods(cellule):
    genre = _attribut(cellule, "office", "value-type")
    if genre in ("float", "percentage", "currency"):
        n = float(_attribut(cellule, "office", "value", "0"))
        return int(n) if n.is_integer() else n
    if genre == "boolean":
        return _attribut(cellule, "office", "boolean-value") == "true"
    if genre == "date":
        brut = _attribut(cellule, "office", "date-value", "")
        try:
            return datetime.fromisoformat(brut)
        except ValueError:
            return brut
    texte = "\n".join("".join(p.itertext()) for p in cellule.iter(f"{{{ODS['text']}}}p"))
    return texte or None


def _feuilles_ods(source: Path) -> list[tuple[str, list[list]]]:
    with zipfile.ZipFile(source) as archive:
        if archive.getinfo("content.xml").file_size > TAILLE_MAX_CONVERSION:
            raise ErreurOutil("Ce tableau est trop volumineux.")
        racine = ElementTree.fromstring(archive.read("content.xml"))
    feuilles = []
    for table in racine.iter(f"{{{ODS['table']}}}table"):
        lignes: list[list] = []
        for rangee in table.iter(f"{{{ODS['table']}}}table-row"):
            ligne: list = []
            for cellule in rangee:
                if cellule.tag.rsplit("}", 1)[-1] not in ("table-cell", "covered-table-cell"):
                    continue
                # LibreOffice résume les cellules identiques (souvent des milliers de cellules vides en fin de ligne)
                repetition = int(_attribut(cellule, "table", "number-columns-repeated", "1"))
                ligne += [_valeur_ods(cellule)] * min(repetition, MAX_COLONNES_CONVERSION - len(ligne))
            repetition = int(_attribut(rangee, "table", "number-rows-repeated", "1"))
            lignes += [ligne] * min(repetition, MAX_LIGNES_CONVERSION - len(lignes))
            if len(lignes) >= MAX_LIGNES_CONVERSION:
                break
        feuilles.append((_attribut(table, "table", "name", "Feuille"), _sans_vides(lignes)))
    return feuilles


def _feuilles_csv(source: Path) -> list[tuple[str, list[list]]]:
    brut = source.read_bytes()[:TAILLE_MAX_CONVERSION]
    try:
        texte = brut.decode("utf-8-sig")
    except UnicodeDecodeError:
        texte = brut.decode("cp1252", errors="replace")
    try:
        dialecte = csv.Sniffer().sniff(texte[:5000], delimiters=";,\t")
    except csv.Error:
        dialecte = csv.excel
    lignes = []
    for ligne in csv.reader(io.StringIO(texte), dialecte):
        lignes.append([_nombre(v) if v.strip() else None for v in ligne[:MAX_COLONNES_CONVERSION]])
        if len(lignes) >= MAX_LIGNES_CONVERSION:
            break
    return [(source.stem, _sans_vides(lignes))]


def _classeur(source: Path, valeurs_seules: bool = False) -> Workbook:
    """Classeur openpyxl quel que soit le format : .xlsx tel quel, .xls / .ods / .csv convertis en mémoire."""
    lecteurs = {".xls": _feuilles_xls, ".ods": _feuilles_ods, ".csv": _feuilles_csv}
    try:
        if source.suffix.lower() in EXTENSIONS:
            return load_workbook(source, data_only=valeurs_seules)
        feuilles = lecteurs[source.suffix.lower()](source)
    except ErreurOutil:
        raise
    except Exception as e:  # fichier corrompu, protégé…
        raise ErreurOutil(f"Je n'arrive pas à ouvrir ce tableau ({e}).") from e
    wb = Workbook()
    wb.remove(wb.active)
    for nom, lignes in feuilles:
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", " ", nom)[:31] or "Feuille")
        for ligne in lignes:
            ws.append(ligne)
            for cellule in ws[ws.max_row]:
                if cellule.data_type == "f":  # un texte commençant par « = » n'est pas une formule ici
                    cellule.data_type = "s"
    if not wb.worksheets:
        wb.create_sheet("Feuille")
    return wb


def _destination(source: Path) -> Path:
    """Copie modifiée, toujours au format .xlsx (un .xls ou un .ods est converti)."""
    return copie_modifiee(source if source.suffix.lower() in EXTENSIONS else source.with_suffix(".xlsx"))


def _charger(chemin: str):
    source = fichier_existant(chemin, EXTENSIONS_LECTURE)
    return source, _classeur(source)


# ---------- outils ----------

def creer_tableau_excel(nom: str, colonnes: list[str], lignes: list[list], ligne_total: bool = False) -> str:
    if not colonnes:
        raise ErreurOutil("Il faut au moins une colonne.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Tableau"
    ws.append([str(c) for c in colonnes])
    for ligne in lignes or []:
        ws.append([_nombre(v) for v in list(ligne)[: len(colonnes)]])
    if ligne_total:
        _ecrire_ligne_total(ws)
    _mettre_en_forme(ws)
    chemin = nouveau_fichier(nom, ".xlsx")
    wb.save(chemin)
    ouvrir(chemin)
    return (f"Tableau créé avec {len(lignes or [])} ligne(s) et {len(colonnes)} colonne(s). "
            + decrire_fichier(chemin))


MAX_FEUILLES_LUES = 3
LIGNE_TOTAL = re.compile(r"^\s*(sous[- ]?)?tota(l|ux)\b", re.I)


def _rempli(v) -> bool:
    return v is not None and v != ""


def _texte(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%d/%m/%Y %H:%M")
    return str(v)


def _est_total(ligne: list) -> bool:
    premiere = next((v for v in ligne if _rempli(v)), None)
    return isinstance(premiere, str) and bool(LIGNE_TOTAL.match(premiere))


def _resumer_feuille(ws: Worksheet) -> list[str]:
    """Aperçu et calculs d'une feuille, même si le tableau ne commence pas en A1 (titre, lignes ou colonnes
    vides, ligne Total)."""
    lignes = [list(l) for l in ws.iter_rows(values_only=True)]
    utilisees = [j for j in range(ws.max_column) if any(j < len(l) and _rempli(l[j]) for l in lignes)]
    tableau = [[l[j] if j < len(l) else None for j in utilisees] for l in lignes]
    tableau = [l for l in tableau if any(_rempli(v) for v in l)]
    if not tableau:
        return [f"\n## Feuille « {ws.title} » : vide"]
    # Ligne des titres de colonnes : la première assez remplie (au-dessus, il n'y a que des titres de page)
    largeur = max(sum(map(_rempli, l)) for l in tableau)
    i_entete = next(i for i, l in enumerate(tableau) if sum(map(_rempli, l)) >= min(largeur, max(2, (largeur + 1) // 2)))
    titres = [_texte(v) for l in tableau[:i_entete] for v in l if _rempli(v)]
    entetes = [_texte(v).strip() if _rempli(v) else get_column_letter(utilisees[k] + 1)
               for k, v in enumerate(tableau[i_entete])]
    donnees = [l for l in tableau[i_entete + 1:] if not _est_total(l)]
    totaux = [l for l in tableau[i_entete + 1:] if _est_total(l)]

    parties = [f"\n## Feuille « {ws.title} » : {len(donnees)} ligne(s) de données"]
    if titres:
        parties.append("Titre : " + " — ".join(titres))
    parties.append("Colonnes : " + " | ".join(entetes))
    parties += [" | ".join(_texte(v) for v in l) for l in donnees[:LIGNES_APERCU]]
    if len(donnees) > LIGNES_APERCU:
        parties.append(f"… ({len(donnees) - LIGNES_APERCU} autres lignes non affichées)")
    if totaux:
        parties.append("Ligne(s) de total déjà présentes dans le fichier : "
                       + " ; ".join(" | ".join(_texte(v) for v in l) for l in totaux))
    stats = []
    for k, nom in enumerate(entetes):
        valeurs = [l[k] for l in donnees if _rempli(l[k])]
        nombres = [v for v in valeurs if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if nombres and len(nombres) >= len(valeurs) / 2:
            stats.append(f"- {nom} : total = {round(sum(nombres), 2)}, moyenne = {round(sum(nombres) / len(nombres), 2)}, "
                         f"min = {min(nombres)}, max = {max(nombres)}, nombre de valeurs = {len(nombres)}")
    if stats:
        parties.append("Calculs exacts sur les lignes de données, sans les lignes de total "
                       "(faits par le programme, à utiliser tels quels) :")
        parties += stats
    return parties


def lire_tableau_excel(chemin: str, dossier: str = "", feuille: str = "") -> str:
    source = fichier_existant(chemin, EXTENSIONS_LECTURE, dossier)
    wb = _classeur(source, valeurs_seules=True)
    feuilles = wb.worksheets
    if feuille:
        cle = normaliser(feuille)
        feuilles = [ws for ws in wb.worksheets if normaliser(ws.title) == cle] or \
                   [ws for ws in wb.worksheets if cle in normaliser(ws.title)]
        if not feuilles:
            raise ErreurOutil(f"Il n'y a pas de feuille « {feuille} ». Feuilles : {', '.join(wb.sheetnames)}.")
    parties = [f"Fichier : {source.name} (emplacement : {emplacement_lisible(source)}). Chemin pour les outils : {source}",
               f"Feuilles : {', '.join(wb.sheetnames)}"]
    for ws in feuilles[:MAX_FEUILLES_LUES]:
        parties += _resumer_feuille(ws)
    if len(feuilles) > MAX_FEUILLES_LUES:
        autres = ", ".join(f"« {ws.title} »" for ws in feuilles[MAX_FEUILLES_LUES:])
        parties.append(f"\nFeuilles non lues : {autres}. Pour en lire une, rappelle l'outil avec feuille=son nom.")
    return "\n".join(parties)


TYPES_MODIFICATION = ("ajouter_ligne", "ajouter_colonne", "ajouter_total", "trier", "modifier_cellule",
                      "supprimer_ligne", "supprimer_colonne", "mettre_en_forme")
SYNONYMES_MODIFICATION = {"tri": "trier", "trie": "trier", "ligne": "ajouter_ligne", "ajout_ligne": "ajouter_ligne",
                          "colonne": "ajouter_colonne", "ajout_colonne": "ajouter_colonne", "total": "ajouter_total",
                          "modifier": "modifier_cellule", "cellule": "modifier_cellule", "supprimer": "supprimer_ligne", "effacer_ligne": "supprimer_ligne",
                          "effacer_colonne": "supprimer_colonne",
                          "mise_en_forme": "mettre_en_forme", "mettre_en_page": "mettre_en_forme"}


def _modification(action) -> dict:
    """Le petit modèle écrit parfois {"ajouter_ligne": {...}} ou {"type": "tri"} : on remet la forme attendue."""
    if isinstance(action, str):
        action = {"type": action}
    if not isinstance(action, dict):
        action = {}
    type_ = str(action.get("type") or action.get("action") or "").lower().replace(" ", "_")
    if not type_:
        cle = next((c for c in action if SYNONYMES_MODIFICATION.get(c.lower(), c.lower()) in TYPES_MODIFICATION), None)
        if cle is not None:
            type_, contenu = cle.lower(), action[cle]
            action = {**{k: v for k, v in action.items() if k != cle}, **(contenu if isinstance(contenu, dict) else {})}
    return {**action, "type": SYNONYMES_MODIFICATION.get(type_, type_)}


def modifier_tableau_excel(chemin: str, actions: list[dict]) -> str:
    source, wb = _charger(chemin)
    ws = wb.active
    faits = []
    for action in map(_modification, actions or []):
        type_ = action["type"]
        if type_ == "ajouter_ligne":
            valeurs = [_nombre(v) for v in action.get("valeurs", [])]
            total = _a_ligne_total(ws)
            if total:
                ws.insert_rows(ws.max_row)
                for i, v in enumerate(valeurs, start=1):
                    ws.cell(row=ws.max_row - 1, column=i, value=v)
                _ecrire_ligne_total(ws)
            else:
                ws.append(valeurs)
            faits.append("ligne ajoutée")
        elif type_ == "ajouter_colonne":
            nom = str(action.get("nom", "Nouvelle colonne"))
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=nom)
            operation = str(action.get("operation", "")).lower()
            if operation:
                symboles = {"additionner": "+", "soustraire": "-", "multiplier": "*", "diviser": "/"}
                if operation not in symboles:
                    raise ErreurOutil("Opération inconnue : utilisez additionner, soustraire, multiplier ou diviser.")
                noms = list(action.get("colonnes") or [])
                # « divise le montant par 12 » : une colonne et un nombre (champ nombre, ou nombre mis dans colonnes)
                nombre = next((action[c] for c in ("nombre", "valeur", "par") if action.get(c) not in (None, "")), None)
                if nombre is None and len(noms) == 2 and isinstance(_nombre(noms[1]), (int, float)):
                    nombre = noms.pop()
                if nombre is not None:
                    if len(noms) != 1 or not isinstance(_nombre(nombre), (int, float)):
                        raise ErreurOutil("Pour calculer avec un nombre, indique une seule colonne et le nombre "
                                          '(ex. {"operation": "diviser", "colonnes": ["Montant"], "nombre": 12}).')
                    if operation == "diviser" and _nombre(nombre) == 0:
                        raise ErreurOutil("Impossible de diviser par zéro.")
                    lettre = get_column_letter(_index_colonne(ws, noms[0]))
                    operandes = lambda r: (f"{lettre}{r}", str(_nombre(nombre)))
                elif len(noms) == 2:
                    a, b = (get_column_letter(_index_colonne(ws, n)) for n in noms)
                    operandes = lambda r: (f"{a}{r}", f"{b}{r}")
                else:
                    raise ErreurOutil("Pour un calcul, indique deux colonnes, ou une colonne et un nombre "
                                      '(ex. {"operation": "diviser", "colonnes": ["Montant"], "nombre": 12}).')
                for r in range(2, _derniere_ligne_donnees(ws) + 1):
                    gauche, droite = operandes(r)
                    ws.cell(row=r, column=col, value=f"={gauche}{symboles[operation]}{droite}")
            else:
                for r, v in enumerate(action.get("valeurs", []), start=2):
                    ws.cell(row=r, column=col, value=_nombre(v))
            if _a_ligne_total(ws) and operation:
                _ecrire_ligne_total(ws)
            faits.append(f"colonne « {nom} » ajoutée")
        elif type_ == "ajouter_total":
            noms = action.get("colonnes") or []
            _ecrire_ligne_total(ws, [_index_colonne(ws, n) for n in noms] or None)
            faits.append("ligne de total ajoutée")
        elif type_ == "trier":
            col = _index_colonne(ws, action.get("colonne", ""))
            decroissant = str(action.get("ordre", "croissant")).lower().startswith("d")
            fin = _derniere_ligne_donnees(ws)
            lignes = [(r, [c.value for c in ligne]) for r, ligne in enumerate(ws.iter_rows(min_row=2, max_row=fin), start=2)]

            def cle(item):
                v = item[1][col - 1]
                return (0, v, "") if isinstance(v, (int, float)) else (1, 0, str(v or "").lower())

            lignes.sort(key=cle, reverse=decroissant)
            for r, (ancienne, ligne) in enumerate(lignes, start=2):
                for c, v in enumerate(ligne, start=1):
                    if isinstance(v, str) and v.startswith("=") and ancienne != r:
                        lettre = get_column_letter(c)
                        v = Translator(v, origin=f"{lettre}{ancienne}").translate_formula(f"{lettre}{r}")
                    ws.cell(row=r, column=c, value=v)
            faits.append(f"tableau trié par « {ws.cell(row=1, column=col).value} »")
        elif type_ == "modifier_cellule":
            ligne = int(action.get("ligne", 0))
            if ligne < 1 or ligne > _derniere_ligne_donnees(ws) - 1:
                raise ErreurOutil(f"La ligne {ligne} n'existe pas (le tableau a {_derniere_ligne_donnees(ws) - 1} lignes).")
            col = _index_colonne(ws, action.get("colonne", ""))
            ws.cell(row=ligne + 1, column=col, value=_nombre(action.get("valeur")))
            faits.append(f"cellule ligne {ligne}, colonne « {ws.cell(row=1, column=col).value} » modifiée")
        elif type_ == "supprimer_colonne":
            col = _index_colonne(ws, action.get("colonne", ""))
            nom = ws.cell(row=1, column=col).value
            ws.delete_cols(col)
            faits.append(f"colonne « {nom} » supprimée")
        elif type_ == "supprimer_ligne":
            ligne = int(action.get("ligne", 0))
            if ligne < 1 or ligne > _derniere_ligne_donnees(ws) - 1:
                raise ErreurOutil(f"La ligne {ligne} n'existe pas.")
            ws.delete_rows(ligne + 1)
            if _a_ligne_total(ws):
                _ecrire_ligne_total(ws)
            faits.append(f"ligne {ligne} supprimée")
        elif type_ == "mettre_en_forme":
            faits.append("mise en forme appliquée")
        else:
            raise ErreurOutil(f"Modification inconnue : « {type_} ». Chaque modification a un champ type parmi "
                              f"{', '.join(TYPES_MODIFICATION)}. Ex : "
                              '[{"type": "ajouter_ligne", "valeurs": ["Électricité", 95]}]. ' "Rien n'a été modifié.")
    if not faits:
        raise ErreurOutil("Aucune modification demandée.")
    _mettre_en_forme(ws)
    destination = _destination(source)
    wb.save(destination)
    ouvrir(destination)
    return (f"Modifications faites : {', '.join(faits)}. L'original n'a pas été touché. "
            "Nouveau fichier. " + decrire_fichier(destination))


def ajouter_graphique_excel(chemin: str, type: str, colonne_x: str, colonnes_y: list[str], titre: str = "") -> str:
    source, wb = _charger(chemin)
    ws = wb.active
    fin = _derniere_ligne_donnees(ws)
    if fin < 2:
        raise ErreurOutil("Le tableau ne contient pas de données pour faire un graphique.")
    type_ = (type or "barres").lower()
    if type_.startswith("cam") or type_ == "secteurs":
        graphique = PieChart()
        colonnes_y = colonnes_y[:1]
    elif type_.startswith("cour") or type_.startswith("lign"):
        graphique = LineChart()
    else:
        graphique = BarChart()
    graphique.title = titre or None
    graphique.height, graphique.width = 9, 16
    col_x = _index_colonne(ws, colonne_x)
    for nom in colonnes_y or []:
        col = _index_colonne(ws, nom)
        graphique.add_data(Reference(ws, min_col=col, min_row=1, max_row=fin), titles_from_data=True)
    if not graphique.series:
        raise ErreurOutil("Indiquez au moins une colonne de valeurs pour le graphique.")
    graphique.set_categories(Reference(ws, min_col=col_x, min_row=2, max_row=fin))
    ws.add_chart(graphique, f"{get_column_letter(ws.max_column + 2)}2")
    destination = _destination(source)
    wb.save(destination)
    ouvrir(destination)
    return "Graphique ajouté à droite du tableau (copie du fichier). " + decrire_fichier(destination)


def tableau(action: str = "", chemin: str = "", nom: str = "", colonnes: list[str] | None = None,
            lignes: list[list] | None = None, ligne_total: bool = False, dossier: str = "", feuille: str = "",
            modifications: list[dict] | None = None, graphique: str = "", colonne_x: str = "",
            colonnes_y: list[str] | None = None, titre: str = "", **autres) -> str:
    """Outil unique pour les tableaux : un petit modèle choisit mieux parmi moins d'outils."""
    # Noms de paramètres des anciens outils (raccourcis enregistrés, ou modèle qui s'en souvient)
    modifications = modifications or autres.get("actions")
    graphique = graphique or autres.get("type", "")
    action = (action or "").lower()
    if action not in ACTIONS_TABLEAU:  # action oubliée : on la devine d'après les paramètres donnés
        action = ("creer" if colonnes else "modifier" if modifications else
                  "graphique" if graphique or colonne_x or colonnes_y else "lire" if chemin else "")
    if action == "lire":
        if not chemin:
            raise ErreurOutil("Indique le nom du fichier à lire (chemin).")
        return lire_tableau_excel(chemin, dossier, feuille)
    if action == "creer":
        return creer_tableau_excel(nom or Path(chemin).stem, colonnes or [], lignes or [], ligne_total)
    if action == "modifier":
        if not modifications:
            raise ErreurOutil("Indique les modifications à faire (modifications : liste d'actions avec un champ type).")
        return modifier_tableau_excel(chemin, modifications)
    if action == "graphique":
        return ajouter_graphique_excel(chemin, graphique, colonne_x, colonnes_y or [], titre)
    raise ErreurOutil("Action inconnue : utilise « lire », « creer », « modifier » ou « graphique ».")


ACTIONS_TABLEAU = ("lire", "creer", "modifier", "graphique")
# Anciens outils séparés → action de l'outil tableau
ANCIENS_OUTILS = {"lire_tableau_excel": "lire", "creer_tableau_excel": "creer",
                  "modifier_tableau_excel": "modifier", "ajouter_graphique_excel": "graphique"}

SCHEMAS = [
    {
        "name": "tableau",
        "description": (
            "Tableaux Excel (.xlsx, .xls), LibreOffice Calc (.ods) et CSV. "
            "action=lire : colonnes, premières lignes et totaux/moyennes exacts calculés par le programme ; le fichier "
            "peut être n'importe où (ex. sur le Bureau), une partie de son nom suffit. "
            "action=creer : crée un nouveau fichier Excel (nom, colonnes, lignes) et l'ouvre. "
            "action=modifier : enregistre une copie modifiée (modifications), l'original reste intact. "
            "action=graphique : ajoute un graphique (graphique, colonne_x, colonnes_y) dans une copie du fichier."),
        "parameters": {
            "type": "object",
            "properties": {
                "action": {"type": "string", "enum": list(ACTIONS_TABLEAU)},
                "chemin": {"type": "string", "description": "Fichier à lire, modifier ou compléter : chemin complet ou "
                                                            "nom (ou partie du nom)"},
                "dossier": {"type": "string", "description": "Pour lire : où chercher, seulement si l'utilisateur l'a dit "
                                                             "(Bureau, Documents, Téléchargements…). Sinon, laisser vide"},
                "feuille": {"type": "string", "description": "Pour lire : nom d'une feuille précise (facultatif)"},
                "nom": {"type": "string", "description": "Pour creer : nom du fichier, sans extension. Ex : Dépenses octobre"},
                "colonnes": {"type": "array", "items": {"type": "string"}, "description": "Pour creer : titres des colonnes"},
                "lignes": {"type": "array", "items": {"type": "array", "items": {}},
                           "description": "Pour creer : une liste de valeurs par ligne, dans l'ordre des colonnes. Peut être vide."},
                "ligne_total": {"type": "boolean", "description": "Pour creer : ajouter une ligne Total qui additionne "
                                                                  "les colonnes de nombres"},
                "modifications": {"type": "array", "items": {"type": "object"}, "description": (
                    "Pour modifier : liste d'actions, chacune avec un champ type : ajouter_ligne {valeurs}, "
                    "ajouter_colonne {nom, valeurs} ou {nom, operation, colonnes:[A,B]} ou {nom, operation, colonnes:[A], nombre} "
                    "avec operation = "
                    "additionner|soustraire|multiplier|diviser, ajouter_total {colonnes?}, trier {colonne, ordre: "
                    "croissant|decroissant}, modifier_cellule {ligne, colonne, valeur}, supprimer_ligne {ligne}, supprimer_colonne {colonne}, "
                    "mettre_en_forme {}. Les lignes sont numérotées à partir de 1 sous les titres. "
                    'Ex : [{"type": "ajouter_ligne", "valeurs": ["Électricité", 95]}, '
                    '{"type": "trier", "colonne": "Montant", "ordre": "decroissant"}]')},
                "graphique": {"type": "string", "enum": ["barres", "courbe", "camembert"], "description": "Pour graphique"},
                "colonne_x": {"type": "string", "description": "Pour graphique : colonne des étiquettes (ex : Mois)"},
                "colonnes_y": {"type": "array", "items": {"type": "string"},
                               "description": "Pour graphique : colonnes de valeurs à dessiner"},
                "titre": {"type": "string", "description": "Pour graphique : titre (facultatif)"},
            },
            "required": ["action"],
        },
    },
]

FONCTIONS = {"tableau": tableau}
