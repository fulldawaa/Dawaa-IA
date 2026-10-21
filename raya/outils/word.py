"""Outils documents : créer un fichier Word, lire un Word / PDF / texte."""
import re
import zipfile
from xml.etree import ElementTree

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader

from raya.outils.fichiers import (
    ErreurOutil, decrire_fichier, emplacement_lisible, fichier_existant, nouveau_fichier, ouvrir,
)

LONGUEUR_MAX_LECTURE = 6000
GRAS = re.compile(r"\*\*(.+?)\*\*")


def _ajouter_texte(paragraphe, texte: str) -> None:
    """Écrit le texte en gérant le **gras** façon markdown."""
    position = 0
    for m in GRAS.finditer(texte):
        if m.start() > position:
            paragraphe.add_run(texte[position:m.start()])
        paragraphe.add_run(m.group(1)).bold = True
        position = m.end()
    if position < len(texte):
        paragraphe.add_run(texte[position:])


def creer_document_word(titre: str, contenu: str, nom_fichier: str = "") -> str:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if titre:
        doc.add_heading(titre, level=1)
    for ligne in (contenu or "").replace("\r\n", "\n").split("\n"):
        brute = ligne.rstrip()
        if not brute.strip():
            continue
        if m := re.match(r"^(#{1,3})\s+(.*)", brute):
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)) + 1, 3))
        elif m := re.match(r"^\s*[-*•]\s+(.*)", brute):
            _ajouter_texte(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", brute):
            _ajouter_texte(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            _ajouter_texte(doc.add_paragraph(), brute)
    chemin = nouveau_fichier(nom_fichier or titre or "Document", ".docx")
    doc.save(chemin)
    ouvrir(chemin)
    return "Document Word créé. " + decrire_fichier(chemin)


ODF = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
}
TAILLE_MAX_ODT = 50 * 1024 * 1024


def _balise(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _texte_en_ligne(element) -> str:
    """Texte d'un paragraphe ODF, avec ses espaces, tabulations et retours à la ligne."""
    morceaux = [element.text or ""]
    for enfant in element:
        balise = _balise(enfant)
        if balise == "s":
            morceaux.append(" " * int(enfant.get(f"{{{ODF['text']}}}c", "1")))
        elif balise == "tab":
            morceaux.append("\t")
        elif balise == "line-break":
            morceaux.append("\n")
        elif balise not in ("note", "annotation", "bookmark", "bookmark-start", "bookmark-end"):
            morceaux.append(_texte_en_ligne(enfant))
        morceaux.append(enfant.tail or "")
    return "".join(morceaux)


def _texte_odt(chemin) -> str:
    """Texte d'un document LibreOffice / OpenOffice (.odt) : une archive ZIP contenant content.xml."""
    with zipfile.ZipFile(chemin) as archive:
        if archive.getinfo("content.xml").file_size > TAILLE_MAX_ODT:
            raise ErreurOutil("Ce document est trop volumineux.")
        racine = ElementTree.fromstring(archive.read("content.xml"))
    corps = racine.find(f".//{{{ODF['office']}}}text")
    lignes: list[str] = []

    def parcourir(element, puce: str = "") -> None:
        for enfant in element:
            balise = _balise(enfant)
            if balise in ("p", "h"):
                lignes.append(puce + _texte_en_ligne(enfant))
            elif balise == "list-item":
                parcourir(enfant, puce="- ")
            elif balise == "table":
                for ligne in enfant.iter(f"{{{ODF['table']}}}table-row"):
                    cellules = [" ".join(_texte_en_ligne(p) for p in cellule.iter(f"{{{ODF['text']}}}p")).strip()
                                for cellule in ligne if _balise(cellule) == "table-cell"]
                    lignes.append(" | ".join(cellules))
            elif balise not in ("tracked-changes", "sequence-decls", "variable-decls"):
                parcourir(enfant, puce)

    if corps is not None:
        parcourir(corps)
    return "\n".join(l for l in lignes if l.strip())


def lire_document(chemin: str, dossier: str = "") -> str:
    source = fichier_existant(chemin, (".docx", ".odt", ".pdf", ".txt", ".md"), dossier)
    extension = source.suffix.lower()
    try:
        if extension == ".odt":
            texte = _texte_odt(source)
        elif extension == ".docx":
            doc = Document(source)
            texte = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                for ligne in table.rows:
                    texte += "\n" + " | ".join(c.text.strip() for c in ligne.cells)
        elif extension == ".pdf":
            lecteur = PdfReader(source)
            texte = "\n".join((page.extract_text() or "") for page in lecteur.pages)
        else:
            texte = source.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        raise ErreurOutil(f"Je n'arrive pas à lire ce document ({e}).") from e
    texte = re.sub(r"\n{3,}", "\n\n", texte).strip()
    if not texte:
        return "Le document semble vide (ou c'est un PDF scanné sans texte, que je ne sais pas lire)."
    suite = ""
    if len(texte) > LONGUEUR_MAX_LECTURE:
        suite = f"\n\n[… document tronqué : seuls les {LONGUEUR_MAX_LECTURE} premiers caractères sont affichés sur {len(texte)}]"
        texte = texte[:LONGUEUR_MAX_LECTURE]
    return f"Contenu de {source.name} (emplacement : {emplacement_lisible(source)}) :\n\n{texte}{suite}"


SCHEMAS = [
    {
        "name": "creer_document_word",
        "description": ("Crée un fichier Word (.docx) et l'ouvre. À utiliser quand l'utilisateur veut un fichier/document Word. "
                        "Le contenu peut utiliser '# Titre', '- puce' et **gras**."),
        "parameters": {
            "type": "object",
            "properties": {
                "titre": {"type": "string", "description": "Titre affiché en haut du document (peut être vide pour une lettre)"},
                "contenu": {"type": "string", "description": "Texte complet du document"},
                "nom_fichier": {"type": "string", "description": "Nom court du fichier, sans extension"},
            },
            "required": ["titre", "contenu"],
        },
    },
    {
        "name": "lire_document",
        "description": ("Lit le texte d'un document Word (.docx), LibreOffice (.odt), PDF ou texte, pour le résumer ou "
                        "répondre à des questions. Le fichier peut être n'importe où (ex. sur le Bureau) : "
                        "il suffit d'une partie de son nom."),
        "parameters": {
            "type": "object",
            "properties": {
                "chemin": {"type": "string", "description": "Chemin complet ou nom (ou partie du nom) du fichier"},
                "dossier": {"type": "string", "description": "Où chercher, seulement si l'utilisateur l'a dit : Bureau, Documents, Téléchargements… Sinon, laisser vide"},
            },
            "required": ["chemin"],
        },
    },
]

FONCTIONS = {"creer_document_word": creer_document_word, "lire_document": lire_document}
